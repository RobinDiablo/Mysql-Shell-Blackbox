import pymysql
import csv
import sys
import psutil
import re
import os
import warnings
from datetime import timedelta,datetime
from openpyxl import load_workbook

if not sys.warnoptions:
    warnings.simplefilter("ignore")



wb = load_workbook(r'C:\Users\Balaji Vyshnavy\Desktop\state2.xlsx')
sheet = wb.active

hostname = 'localhost'
username = 'root'
password = 'rootpassword'
database = 'custom1'

conn = pymysql.connect( host=hostname, user=username, passwd=password, db=database )
cur = conn.cursor()

def colu(rez , file ,tn):
    #print(tn)
    with open(f'C:/Users/balaji Vyshnavy/Desktop/evidence/{file}', mode='a') as csv_file:
        fieldnames = rez
        writer = csv.DictWriter(csv_file, fieldnames=tn, )
        writer.writeheader()
        writer =csv.DictWriter(csv_file, fieldnames=fieldnames,)
        writer.writeheader()

def data(file,n):
    test_string=filter(n)
    test_string=int(test_string)
    #print(test_string)

    with open(f'C:/Users/balaji Vyshnavy/Desktop/evidence/{file}', mode='a') as csv_file:
        for i in range(0, test_string):
            pez = cur.fetchone()
            fieldnames = pez
            writer = csv.DictWriter(csv_file, fieldnames=fieldnames)
            writer.writeheader()

def filter(x):
    bad_chars = ['(', ')']
    bc = [',,']
    test_string = x
    test_string = str(test_string)
    for i in bad_chars:
        test_string = test_string.replace(i, '')
    for i in bc:
        test_string = test_string.replace(i, '')
    #print("Resultant list is : " + test_string)
    return test_string
def command_gen(rez):
    bad_chars = ['(', ')', "'"]
    test_string = rez[0]
    test_string = str(test_string)
    #for i in bad_chars:
    #    test_string = test_string.replace(i, '')
    #test_string = test_string.replace(',', '')
    #cmd1='use '+ pez
    cmd2='checksum table '+ test_string
    #print(cmd1)
    return cmd2

def db_gen(rez):
    test_string = rez[1]
    test_string = str(test_string)
    cmd1='use '+ test_string
    return cmd1

with open(f'C:/Users/balaji Vyshnavy/Desktop/evidence/testing.csv', mode='w') as f:
    f.truncate()
    f.close()


def writecsv(hash):
    with open(f'C:/Users/balaji Vyshnavy/Desktop/evidence/testing.csv', mode='a') as csv_file:
        fieldnames = hash
        writer = csv.DictWriter(csv_file, fieldnames=hash, )
        writer.writeheader()
   

cur.execute('select table_name,table_schema from information_schema.tables')
gg=cur.fetchall()
for res in gg:
    cmd2 = command_gen(res)
    cmd1 = db_gen(res)
    #print(cmd1)
    #print(cmd2)
    cur.execute(cmd1)
    cur.execute(cmd2)
    hsh = cur.fetchall()
    writecsv(hsh)



def update(value):
    with open(f'C:/Users/balaji Vyshnavy/Desktop/evidence/logbook.csv', mode='a') as csv_file:
        fieldnames = value
        writer = csv.DictWriter(csv_file, fieldnames=value, )
        writer.writeheader()     

def updatetime(time):
    with open(f'C:/Users/balaji Vyshnavy/Desktop/evidence/logbook.csv', mode='a') as csv_file:
        fieldnames = time
        writer = csv.DictWriter(csv_file, fieldnames=time, )
        writer.writeheader()  

def comparator(file1,file2):
    e=open(f'C:/Users/balaji Vyshnavy/Desktop/evidence/{file1}.csv', mode='r')
    #t2=open(f'C:/Users/balaji Vyshnavy/Desktop/evidence/testing.csv', mode='r')
    f=open(f'C:/Users/balaji Vyshnavy/Desktop/evidence/{file2}.csv', 'r')
    countf=0 
    counte=0 
    rowse = list(csv.reader(e))
    rowsf= list(csv.reader(f))
    dateTimeObj = datetime.now()
    timenow= dateTimeObj.strftime("%Y/%d/%b--%H:%M:%S")
    print("updating")
    print(timenow)
    for i in rowsf:
        countf=countf+1
    for i in rowse:
        counte=counte+1
    if counte>countf:
        count=countf
    else:
        count=counte

    i=0
    
    updatetime(timenow)
    while i < count:
        if rowsf[i]!=rowse[i]:
                update(rowsf[i])
                update(rowse[i])
        i=i+1
    
def getfiles():
    cur.execute('use information_schema')
    cur.execute('select column_name from information_schema.columns where table_name=" columns"')
    rez=cur.fetchall()
    tn="columns"
    colu(rez , 'columns_temp.csv',tn)
    cur.execute('select count(*) from columns')
    n=cur.fetchall()
    cur.execute('select * from columns')
    data('columns_temp.csv',n)
    cur.execute('select column_name from information_schema.columns where table_name=" tables"')
    rez=cur.fetchall()
    tn="tables"
    colu(rez , 'tables_temp.csv',tn)
    cur.execute('select count(*) from tables')
    n=cur.fetchall()
    cur.execute('select * from tables')
    data('tables_temp.csv',n)
    #comparator('columns','columns_temp')




def find():
    dateTimeObj = datetime.now()
    timenow= dateTimeObj.strftime("%H:%M:%S")
    x,y,z=timenow.split(':')
    x=int(x)
    y=int(y)
    z=int(z)
    t1 = timedelta(hours=x, minutes=y,seconds=z)
    t2 = timedelta(hours=0, minutes=20,seconds=00)
    for p in psutil.process_iter():
        line=p.name
        pid=p.pid
        line = str(line)
        line=line[49:-2]
        bad_chars = [", name='", "', started='"]
        bc=["'"]
        test_string = line
        test_string = str(test_string)
        for i in bad_chars:
            test_string = test_string.replace(i, ';')
        for i in bc:
            test_string = test_string.replace(i, '')
        time=test_string[-8:]
        y= test_string.rpartition(';')
        y=y[0]
        name = "".join(re.split("[^a-zA-Z]*", y))
        name=str(name)
        msg='mysqlshexe'
        if name==msg :
            return 1
            break
    return 0
           
        
x=find()     
print(x)

if  x==1:
    print("running")
    sheet['A1'] = 'unsaved'
    wb.save(r'C:\Users\Balaji Vyshnavy\Desktop\state2.xlsx')

elif sheet['A1'].value == 'unsaved':
    
    print("saving")
    sheet['A1'] = 'saved'
    dateTimeObj = datetime.utcnow()
    cal={"Jan":"01","Feb":"02","Mar":"03","Apr":"04","May":"05","Jun":"06","Jul":"07","Aug":"08","Sep":"09","Oct":"10","Nov":"11","Dec":"12"}
    m=dateTimeObj.strftime("%b")
    for x, y in cal.items():
        if m==x:
            m=y
    timenow=dateTimeObj.strftime("%Y-")+str(m)+dateTimeObj.strftime("-%dT%H:%M:%S.%f")
    sheet['B1'] = timenow
    getfiles()
    comparator('tables','tables_temp')
    comparator('checksum','testing')
    wb.save(r'C:\Users\Balaji Vyshnavy\Desktop\state2.xlsx')


conn.close()
















